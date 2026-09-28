"""Native editable Word tables and lossless gallery PDF; bundled document Python."""
import json
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

OUT=Path(__file__).resolve().parents[1]/'runs/paper_comparison'


def border(cell, edge, width):
    tcpr=cell._tc.get_or_add_tcPr()
    borders=tcpr.find(qn('w:tcBorders'))
    if borders is None:
        borders=OxmlElement('w:tcBorders');tcpr.append(borders)
    element=borders.find(qn('w:'+edge))
    if element is None:
        element=OxmlElement('w:'+edge);borders.append(element)
    element.set(qn('w:val'),'single' if width else 'nil')
    element.set(qn('w:sz'),str(width))
    element.set(qn('w:color'),'000000')


def add_table(doc,headers,names,values,digits):
    table=doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths=[5.1]+[4.1]*(len(headers)-1)
    for column,width in zip(table.columns,widths):column.width=Cm(width)
    maxima=[max(row[j] for row in values) for j in range(len(values[0]))]
    rows=[headers]+[[name]+[f'{v:.{digits}f}' for v in row] for name,row in zip(names,values)]
    for i,items in enumerate(rows):
        cells=table.rows[0].cells if i==0 else table.add_row().cells
        for j,(cell,content) in enumerate(zip(cells,items)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for side in ['top','bottom','left','right']:border(cell,side,0)
            if i==0:
                border(cell,'top',16);border(cell,'bottom',8)
            if i==len(rows)-1:border(cell,'bottom',8)
            paragraph=cell.paragraphs[0];paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before=Pt(8);paragraph.paragraph_format.space_after=Pt(8)
            run=paragraph.add_run(content);run.font.size=Pt(12)
            run.bold=i==0 or (j>0 and values[i-1][j-1]==maxima[j-1])
    return table


def document(data,historical=False):
    doc=Document();sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Cm(29.7);sec.page_height=Cm(21)
    sec.top_margin=sec.bottom_margin=Cm(1.3);sec.left_margin=sec.right_margin=Cm(1.6)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    prefix='historical_' if historical else ''
    names=data[prefix+'models'];overall=data[prefix+'overall'];iou=data[prefix+'class_iou']
    doc.add_paragraph('历史实验结果（单独列示，未重新运行）' if historical else '表1  不同方法的分割性能比较')
    add_table(doc,['Model Name','PixAcc','mIoU','F1' if historical else 'F1 (macro)'],names,overall,5)
    doc.add_paragraph('表2  各类别 IoU')
    headers=['Model Name']+[name if historical else f'{name}\n({p*100:.2f}%)' for name,p in zip(data['class_names'],data['class_proportions'])]
    add_table(doc,headers,names,iou,4)
    if historical:
        doc.add_paragraph('注：数值来自原始课程报告，按历史记录保留。训练环境、评估分辨率与新实验不同，不应与新实验结果做直接归因比较。历史F1口径未进一步验证。每列最优值加粗，数值范围0–1。')
    else:
        doc.add_paragraph('注：原始320/80训练/验证划分，种子42，实际训练50轮，普通滑窗验证，无TTA。各方法采用各自的骨干、预训练来源及优化配置；每列最优值分别加粗。')
        doc.add_paragraph('类别占比按80张验证图的真值像素数计算。F1为五个类别F1的算术平均，PixAcc为全部有效像素的分类正确率；表中数值范围为0–1。Ours (no text)为当前ConvNeXt-Tiny + FPN无文本模型，不是历史MMDroneSeg。')
    path=OUT/('paper_tables_historical.docx' if historical else 'paper_tables_current.docx');doc.save(path)
    loaded=Document(path);assert len(loaded.tables)==2
    assert all(len(t.rows)==len(names)+1 for t in loaded.tables)


def main():
    data=json.loads((OUT/'table_data.json').read_text(encoding='utf-8'))
    document(data);document(data,True)
    path=OUT/'five_image_comparison_native.png'
    with Image.open(path) as im:
        w,h=im.size
        pdf=canvas.Canvas(str(OUT/'five_image_comparison.pdf'),pagesize=(w/4,h/4),pageCompression=1)
        pdf.setTitle('Five-image qualitative comparison — same examples as original report')
        pdf.drawImage(ImageReader(im),0,0,width=w/4,height=h/4)
        pdf.showPage();pdf.save()
    print('Created native Word tables and full-resolution comparison PDF.')


if __name__=='__main__':main()
