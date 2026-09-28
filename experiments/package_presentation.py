"""Package the rendered figures as widescreen slides with speaker notes."""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageOps, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parents[1] / 'runs/classroom_figures'


def main():
    manifest = json.loads((OUT/'manifest.json').read_text(encoding='utf-8'))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    prs.core_properties.title = '无人机俯视图分割 · 实验图表'
    prs.core_properties.subject = '单种子验证集结果，50轮训练预算'
    for page in manifest['pages']:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(OUT/(page['name']+'.png')), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = page['title']+'\n\n'+page['notes']
    path = OUT/'无人机分割_课堂实验图表.pptx'
    prs.save(path)
    check = Presentation(path)
    assert len(check.slides) == len(manifest['pages']) == 8
    assert all(len(slide.shapes) == 1 and slide.has_notes_slide for slide in check.slides)
    # Contact sheet is only a preview; full-resolution figures remain separate.
    sheet=Image.new('RGB',(1600,1960),'#e7edf2')
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',22)
    draw=ImageDraw.Draw(sheet)
    for i,page in enumerate(manifest['pages']):
        image=Image.open(OUT/(page['name']+'.png')).convert('RGB')
        assert image.size == (2880,1620)
        thumb=ImageOps.contain(image,(768,432))
        x,y=16+(i%2)*800,16+(i//2)*490
        sheet.paste(thumb,(x,y))
        draw.text((x+8,y+442),page['title'],font=font,fill='#203247')
    sheet.save(OUT/'图表总览.jpg',quality=92)
    print(path)


if __name__=='__main__':
    main()
